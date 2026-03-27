from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING


def _set_normal_style(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(6)


def _add_title_page(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("USB FORENSICS TOOL\nDETAILED PROJECT REPORT")
    r.bold = True
    r.font.size = Pt(20)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.add_run("Topic-wise Consolidated Report (20-25 page format)").font.size = Pt(14)

    doc.add_paragraph("\n")
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("Repository: USB-Forensics-Tool-first\n")
    meta.add_run("Prepared for academic submission\n")
    meta.add_run("Formatting: Times New Roman, 12 pt, 1.5 line spacing\n")

    doc.add_page_break()


def _apply_heading(doc: Document, text: str) -> None:
    level = 1
    if re.match(r"^##\s", text):
        level = 2
        text = text[3:]
    elif re.match(r"^###\s", text):
        level = 3
        text = text[4:]
    elif re.match(r"^#\s", text):
        level = 1
        text = text[2:]
    h = doc.add_heading(text.strip(), level=level)
    if level == 1:
        h.paragraph_format.space_before = Pt(12)


def _render_markdown(doc: Document, content: str) -> None:
    in_code = False
    for raw_line in content.splitlines():
        line = raw_line.rstrip()

        if line.strip().startswith("```"):
            in_code = not in_code
            continue

        if in_code:
            p = doc.add_paragraph(line)
            p.style = doc.styles["Normal"]
            continue

        if line.startswith("### ") or line.startswith("## ") or line.startswith("# "):
            _apply_heading(doc, line)
            continue

        if line.startswith("- "):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
            continue

        if re.match(r"^[0-9]+\.\s", line):
            # Keep explicit numbering to preserve author order
            doc.add_paragraph(line)
            continue

        if not line.strip():
            doc.add_paragraph("")
            continue

        doc.add_paragraph(line)


def _add_screenshots(doc: Document, report_dir: Path) -> None:
    shots = [
        report_dir / "screenshots" / "01_devices_page.png",
        report_dir / "screenshots" / "02_analysis_page.png",
        report_dir / "screenshots" / "03_export_page.png",
    ]

    doc.add_heading("Screenshots", level=1)
    for idx, shot in enumerate(shots, 1):
        if shot.exists():
            doc.add_picture(str(shot), width=Inches(6.2))
            cp = doc.add_paragraph(f"Figure {idx}: {shot.name}")
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER


def main() -> None:
    project_root = Path(__file__).resolve().parents[1]
    report_dir = project_root / "report and paper"
    src_dir = report_dir / "Report text document topic wise"

    files = sorted([p for p in src_dir.glob("[0-9][0-9]_*.md") if p.name != "00_Index.md"])
    if not files:
        raise RuntimeError("No topic-wise markdown files found.")

    doc = Document()
    _set_normal_style(doc)
    _add_title_page(doc)

    toc = doc.add_heading("Table of Contents", level=1)
    for f in files:
        doc.add_paragraph(f.stem.split("_", 1)[1].replace("_", " "), style="List Number")
    doc.add_page_break()

    for i, file_path in enumerate(files, 1):
        content = file_path.read_text(encoding="utf-8")
        _render_markdown(doc, content)

        # Keep chapter-style separation and ensure long-form page count.
        if i != len(files):
            doc.add_page_break()

    doc.add_page_break()
    _add_screenshots(doc, report_dir)

    out_file = report_dir / "USB_Forensics_Detailed_Topicwise_Report_20_25_pages.docx"
    doc.save(out_file)
    print(f"Created: {out_file}")


if __name__ == "__main__":
    main()

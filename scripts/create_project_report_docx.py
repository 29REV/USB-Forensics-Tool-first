from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
import json

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def add_title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("USB FORENSICS TOOL\nPROJECT REPORT")
    run.bold = True
    run.font.size = Pt(22)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("Comprehensive Technical Documentation and Runtime Analysis")
    r2.font.size = Pt(14)

    doc.add_paragraph("\n")

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("Repository: USB-Forensics-Tool-first\nBranch: ReportAndPaper\n").bold = True
    meta.add_run(f"Report Generated (UTC): {datetime.now(timezone.utc).isoformat()}\n")
    meta.add_run("Platform: Windows\nLanguage: Python 3.13")

    doc.add_page_break()


def add_section_heading(doc: Document, title: str) -> None:
    h = doc.add_heading(title, level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_chapter_heading(doc: Document, chapter_no: int, title: str) -> None:
    add_section_heading(doc, f"CHAPTER {chapter_no}: {title}")


def add_table_of_contents_manual(doc: Document) -> None:
    add_section_heading(doc, "TABLE OF CONTENTS")
    for line in [
        "Abstract",
        "Chapter 1: Introduction",
        "Chapter 2: Technical Background and Related Concepts",
        "Chapter 3: Methodology and Processing Pipeline",
        "Chapter 4: Module-by-Module System Design",
        "Chapter 5: Runtime Execution, Screenshots, and Results",
        "Chapter 6: Security, Limitations, and Future Enhancements",
        "Conclusion",
        "References",
    ]:
        doc.add_paragraph(line, style="List Bullet")
    doc.add_page_break()


def add_abstract(doc: Document) -> None:
    add_section_heading(doc, "ABSTRACT")
    doc.add_paragraph(
        "The USB Forensics Tool is a Python-based digital forensics platform that performs endpoint-level USB artifact analysis through a combination of device enumeration, Windows telemetry processing, anomaly scoring, optional USB trace intelligence, and multi-format reporting. "
        "The project supports both a professional Tkinter GUI workflow and script-driven processing for evidence generation. "
        "Its architecture is modular and designed for graceful degradation: when privileged Windows APIs or optional packages are unavailable, the system continues in reduced mode with fallback logic. "
        "This report documents the complete technical stack, internal methods, algorithms, module responsibilities, runtime behavior, generated evidence artifacts, and practical limitations."
    )
    doc.add_page_break()


def add_tech_stack_table(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=3)
    hdr = table.rows[0].cells
    hdr[0].text = "Layer"
    hdr[1].text = "Technologies"
    hdr[2].text = "Purpose"

    rows = [
        ("Language and Runtime", "Python 3.13", "Primary implementation language"),
        ("GUI", "Tkinter, ttk", "Professional desktop interface and workflows"),
        ("Windows Integration", "pywin32, WMI, winreg, win32evtlog", "USB device and event acquisition"),
        ("Report Export", "csv, json, openpyxl, reportlab", "CSV, JSON, XLSX, and PDF outputs"),
        ("Image and Media", "Pillow (ImageGrab)", "GUI screenshots and visual evidence support"),
        ("Trace Analytics", "Wireshark/tshark JSON bridge", "Optional packet-level USB behavior analysis"),
        ("URB Capture", "ETW, logman, etl-parser", "USB Request Block capture and parsing"),
        ("Extensions", "JSON IPI bridge + pluggable modules", "Recall, firmware validation, AI narrative extensions"),
    ]

    for layer, tech, purpose in rows:
        c = table.add_row().cells
        c[0].text = layer
        c[1].text = tech
        c[2].text = purpose


def add_runtime_summary_table(doc: Document, runtime: dict) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Value"

    metrics = [
        ("Runtime Timestamp (UTC)", str(runtime.get("timestamp_utc", ""))),
        ("Detected Devices", str(runtime.get("device_count", 0))),
        ("Suspicious Devices", str(runtime.get("suspicious_count", 0))),
    ]

    for k, v in metrics:
        cells = table.add_row().cells
        cells[0].text = k
        cells[1].text = v


def add_screenshots(doc: Document, screenshots_dir: Path) -> None:
    add_section_heading(doc, "EXECUTION SCREENSHOTS")

    screenshots = [
        ("01_devices_page.png", "Figure 1: All Connected USB Devices page after runtime scan."),
        ("02_analysis_page.png", "Figure 2: Advanced Analysis page showing selected USB device context."),
        ("03_export_page.png", "Figure 3: Export Reports page with format selection and export action."),
    ]

    for filename, caption in screenshots:
        path = screenshots_dir / filename
        if path.exists():
            doc.add_picture(str(path), width=Inches(6.5))
            p = doc.add_paragraph(caption)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_module_matrix(doc: Document) -> None:
    add_section_heading(doc, "MODULE RESPONSIBILITY MATRIX")
    table = doc.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "Module"
    table.rows[0].cells[1].text = "Primary Methods"
    table.rows[0].cells[2].text = "Responsibility"

    modules = [
        ("core/usb_device_manager.py", "get_all_usb_devices, classify_device_type, extract_vid_pid", "Enumerate and classify USB devices using WMI or PowerShell fallback"),
        ("core/registry_parser.py", "parse_registry", "Acquire USB registry artifacts from USBSTOR with fallback entries"),
        ("core/eventlog_parser.py", "parse_event_logs", "Acquire connect/disconnect events from Windows Event Log"),
        ("core/correlation.py", "correlate", "Unify registry and event sources into DeviceRecord timelines"),
        ("core/analysis.py", "enrich_summary, compute_anomaly_score, detect_suspicious", "Compute forensic metrics, scoring, and suspicious flags"),
        ("utils/report_generator.py", "write_csv, write_json, write_xlsx, write_pdf", "Generate formal forensic outputs in multiple formats"),
        ("core/wireshark_bridge.py", "load_capture, runtime_info", "Load and parse tshark JSON for USB trace packets"),
        ("core/usb_trace_analysis.py", "analyze_usb_trace, match_trace_report", "Detect suspicious communication patterns in packet-level traces"),
        ("core/urb_capture.py", "parse_etl_file, URBCapture", "Capture and parse ETW URB operations for low-level USB telemetry"),
        ("extensions/bridge.py", "IPIBridge, BaseToolExtensionInterface", "Route JSON messages to extension modules without core edits"),
    ]

    for module, methods, responsibility in modules:
        cells = table.add_row().cells
        cells[0].text = module
        cells[1].text = methods
        cells[2].text = responsibility


def add_device_table(doc: Document, runtime: dict) -> None:
    add_section_heading(doc, "RUNTIME DEVICE SNAPSHOT")
    table = doc.add_table(rows=1, cols=5)
    table.rows[0].cells[0].text = "Name"
    table.rows[0].cells[1].text = "Type"
    table.rows[0].cells[2].text = "VID:PID"
    table.rows[0].cells[3].text = "Serial"
    table.rows[0].cells[4].text = "Anomaly"

    for d in runtime.get("devices", []):
        row = table.add_row().cells
        row[0].text = str(d.get("name", ""))
        row[1].text = str(d.get("device_type", ""))
        row[2].text = f"{d.get('vid', '')}:{d.get('pid', '')}"
        row[3].text = str(d.get("serial", ""))
        row[4].text = str(d.get("anomaly_score", ""))


def build_report() -> Path:
    project_root = Path(__file__).resolve().parents[1]
    report_dir = project_root / "report and paper"

    runtime_root = report_dir / "runtime_evidence"
    runs = sorted([p for p in runtime_root.iterdir() if p.is_dir()])
    if not runs:
        raise RuntimeError("No runtime evidence directories found. Run generate_runtime_evidence.py first.")

    latest_run = runs[-1]
    runtime_summary = latest_run / "runtime_summary.json"
    runtime = json.loads(runtime_summary.read_text(encoding="utf-8"))

    screenshots_dir = report_dir / "screenshots"

    doc = Document()

    add_title(doc)
    add_table_of_contents_manual(doc)
    add_abstract(doc)

    add_chapter_heading(doc, 1, "INTRODUCTION")
    doc.add_paragraph(
        "USB media and removable peripherals are frequent vectors in enterprise and personal device incidents. "
        "This project was developed as an end-to-end forensic toolkit to discover USB-connected hardware, reconstruct activity timelines, score anomalies, and export investigation-ready reports."
    )
    doc.add_paragraph(
        "The solution targets Windows environments where registry hives, event logs, and USB stack telemetry provide high forensic value. "
        "The GUI design emphasizes analyst productivity through page-based workflows: Devices, Storage, Timeline, Analysis, URB Capture, Security, and Export."
    )

    add_chapter_heading(doc, 2, "TECHNICAL BACKGROUND AND RELATED CONCEPTS")
    doc.add_paragraph(
        "The architecture combines host artifact analysis (registry and event logs) with optional deep packet/trace telemetry. "
        "A key design choice is graceful degradation: the tool should remain runnable and demonstrable even when pywin32, admin privileges, or trace tooling are unavailable."
    )
    add_tech_stack_table(doc)

    add_chapter_heading(doc, 3, "METHODOLOGY AND PROCESSING PIPELINE")
    for step in [
        "Step 1 - Device acquisition: collect USB hardware entities via WMI, then PowerShell fallback, then mock fallback.",
        "Step 2 - Artifact collection: parse USBSTOR registry keys and relevant Event Log IDs for connection behavior.",
        "Step 3 - Correlation: merge multiple artifact streams into unified DeviceRecord timelines.",
        "Step 4 - Enrichment: derive storage, folder, deleted-file, and reputation analysis per device.",
        "Step 5 - Scoring: compute anomaly score (0-100) using serial presence, volume of events, storage pressure, and optional trace indicators.",
        "Step 6 - Reporting: export enriched findings to CSV, JSON, XLSX, PDF, and detailed JSON.",
        "Step 7 - Optional deep analysis: parse tshark JSON/pcap-derived capture and ETW URB traces for suspicious communication signatures.",
    ]:
        doc.add_paragraph(step, style="List Number")

    add_chapter_heading(doc, 4, "MODULE-BY-MODULE SYSTEM DESIGN")
    add_module_matrix(doc)

    add_chapter_heading(doc, 5, "RUNTIME EXECUTION, SCREENSHOTS, AND RESULTS")
    doc.add_paragraph(
        "The application was executed on Windows, and screenshots were captured programmatically from live GUI windows. "
        "The runtime evidence pipeline generated new report files (CSV/JSON/XLSX/PDF/detailed JSON) and captured the current USB device state."
    )
    add_runtime_summary_table(doc, runtime)
    add_device_table(doc, runtime)
    add_screenshots(doc, screenshots_dir)

    add_section_heading(doc, "GENERATED EVIDENCE ARTIFACTS")
    generated = runtime.get("generated_files", {})
    for key, value in generated.items():
        doc.add_paragraph(f"{key}: {value}", style="List Bullet")

    add_chapter_heading(doc, 6, "SECURITY, LIMITATIONS, AND FUTURE ENHANCEMENTS")
    doc.add_paragraph("Current security strengths:")
    for item in [
        "Cross-source evidence collection (registry plus event timeline plus optional traces).",
        "Heuristic suspicious pattern detection and anomaly scoring.",
        "Optional packet-level behavioral analytics for identifying potential BadUSB style indicators.",
        "Pluggable extension framework for recall, firmware, and AI-assisted narrative analysis.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_paragraph("Current limitations:")
    for item in [
        "Many advanced features require administrator privileges and Windows APIs.",
        "Some enrichment sections depend on external or simulated data sources.",
        "Scoring is heuristic-based and should complement, not replace, analyst judgment.",
        "Trace and URB analyses are optional and depend on capture availability and tooling readiness.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_paragraph("Recommended future work:")
    for item in [
        "Add cryptographic evidence chaining and tamper-evident audit logs.",
        "Introduce signed plugin loading and extension permission controls.",
        "Expand model-driven anomaly scoring using historical baselines.",
        "Automate timeline visualization and case-oriented report templates.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    add_section_heading(doc, "CONCLUSION")
    doc.add_paragraph(
        "The USB Forensics Tool provides a practical and extensible forensic platform for investigators who need rapid USB visibility with optional deep protocol inspection. "
        "Its modular architecture, fallback-safe implementation, and strong reporting support make it suitable for academic demonstration as well as controlled operational workflows."
    )

    add_section_heading(doc, "REFERENCES")
    refs = [
        "Project repository source files: main.py, gui/app.py, core/*.py, utils/report_generator.py",
        "Windows Event Log and Registry artifacts relevant to USB forensic reconstruction",
        "Wireshark/tshark documentation for USB capture parsing",
        "Python package documentation: pywin32, openpyxl, reportlab, Pillow, python-docx",
    ]
    for r in refs:
        doc.add_paragraph(r, style="List Bullet")

    out_path = report_dir / "USB_Forensics_Project_Report.docx"
    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    path = build_report()
    print(f"Report written: {path}")
